import subprocess
import shutil
from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
import re


space_groups = {
    # Triclinic
    1:   r"$P1$",
    2:   r"$P\bar{1}$",

    # Monoclinic
    3:   r"$P2$",
    4:   r"$P2_1$",
    5:   r"$C2$",
    6:   r"$Pm$",
    7:   r"$Pc$",
    8:   r"$Cm$",
    9:   r"$Cc$",
    10:  r"$P2/m$",
    11:  r"$P2_1/m$",
    12:  r"$C2/m$",
    13:  r"$P2/c$",
    14:  r"$P2_1/c$",
    15:  r"$C2/c$",

    # Orthorhombic
    16:  r"$P222$",
    17:  r"$P222_1$",
    18:  r"$P2_{1}2_{1}2$",
    19:  r"$P2_{1}2_{1}2_{1}$",
    20:  r"$C222_1$",
    21:  r"$C222$",
    22:  r"$F222$",
    23:  r"$I222$",
    24:  r"$I2_{1}2_{1}2_{1}$",
    25:  r"$Pmm2$",
    26:  r"$Pmc2_1$",
    27:  r"$Pcc2$",
    28:  r"$Pma2$",
    29:  r"$Pca2_1$",
    30:  r"$Pnc2$",
    31:  r"$Pmn2_1$",
    32:  r"$Pba2$",
    33:  r"$Pna2_1$",
    34:  r"$Pnn2$",
    35:  r"$Cmm2$",
    36:  r"$Cmc2_1$",
    37:  r"$Ccc2$",
    38:  r"$Amm2$",
    39:  r"$Aem2$",
    40:  r"$Ama2$",
    41:  r"$Aea2$",
    42:  r"$Fmm2$",
    43:  r"$Fdd2$",
    44:  r"$Imm2$",
    45:  r"$Iba2$",
    46:  r"$Ima2$",
    47:  r"$Pmmm$",
    48:  r"$Pnnn$",
    49:  r"$Pccm$",
    50:  r"$Pban$",
    51:  r"$Pmma$",
    52:  r"$Pnna$",
    53:  r"$Pmna$",
    54:  r"$Pcca$",
    55:  r"$Pbam$",
    56:  r"$Pccn$",
    57:  r"$Pbcm$",
    58:  r"$Pnnm$",
    59:  r"$Pmmn$",
    60:  r"$Pbcn$",
    61:  r"$Pbca$",
    62:  r"$Pnma$",
    63:  r"$Cmcm$",
    64:  r"$Cmce$",
    65:  r"$Cmmm$",
    66:  r"$Cccm$",
    67:  r"$Cmme$",
    68:  r"$Ccce$",
    69:  r"$Fmmm$",
    70:  r"$Fddd$",
    71:  r"$Immm$",
    72:  r"$Ibam$",
    73:  r"$Ibca$",
    74:  r"$Imma$",

    # Tetragonal
    75:  r"$P4$",
    76:  r"$P4_1$",
    77:  r"$P4_2$",
    78:  r"$P4_3$",
    79:  r"$I4$",
    80:  r"$I4_1$",
    81:  r"$P\bar{4}$",
    82:  r"$I\bar{4}$",
    83:  r"$P4/m$",
    84:  r"$P4_2/m$",
    85:  r"$P4/n$",
    86:  r"$P4_2/n$",
    87:  r"$I4/m$",
    88:  r"$I4_1/a$",
    89:  r"$P422$",
    90:  r"$P42_{1}2$",
    91:  r"$P4_{1}22$",
    92:  r"$P4_{1}2_{1}2$",
    93:  r"$P4_{2}22$",
    94:  r"$P4_{2}2_{1}2$",
    95:  r"$P4_{3}22$",
    96:  r"$P4_{3}2_{1}2$",
    97:  r"$I422$",
    98:  r"$I4_{1}22$",
    99:  r"$P4mm$",
    100: r"$P4bm$",
    101: r"$P4_2cm$",
    102: r"$P4_2nm$",
    103: r"$P4cc$",
    104: r"$P4nc$",
    105: r"$P4_2mc$",
    106: r"$P4_2bc$",
    107: r"$I4mm$",
    108: r"$I4cm$",
    109: r"$I4_1md$",
    110: r"$I4_1cd$",
    111: r"$P\bar{4}2m$",
    112: r"$P\bar{4}2c$",
    113: r"$P\bar{4}2_1m$",
    114: r"$P\bar{4}2_1c$",
    115: r"$P\bar{4}m2$",
    116: r"$P\bar{4}c2$",
    117: r"$P\bar{4}b2$",
    118: r"$P\bar{4}n2$",
    119: r"$I\bar{4}m2$",
    120: r"$I\bar{4}c2$",
    121: r"$I\bar{4}2m$",
    122: r"$I\bar{4}2d$",
    123: r"$P4/mmm$",
    124: r"$P4/mcc$",
    125: r"$P4/nbm$",
    126: r"$P4/nnc$",
    127: r"$P4/mbm$",
    128: r"$P4/mnc$",
    129: r"$P4/nmm$",
    130: r"$P4/ncc$",
    131: r"$P4_2/mmc$",
    132: r"$P4_2/mcm$",
    133: r"$P4_2/nbc$",
    134: r"$P4_2/nnm$",
    135: r"$P4_2/mbc$",
    136: r"$P4_2/mnm$",
    137: r"$P4_2/nmc$",
    138: r"$P4_2/ncm$",
    139: r"$I4/mmm$",
    140: r"$I4/mcm$",
    141: r"$I4_1/amd$",
    142: r"$I4_1/acd$",

    # Trigonal
    143: r"$P3$",
    144: r"$P3_1$",
    145: r"$P3_2$",
    146: r"$R3$",
    147: r"$P\bar{3}$",
    148: r"$R\bar{3}$",
    149: r"$P312$",
    150: r"$P321$",
    151: r"$P3_{1}12$",
    152: r"$P3_{1}21$",
    153: r"$P3_{2}12$",
    154: r"$P3_{2}21$",
    155: r"$R32$",
    156: r"$P3m1$",
    157: r"$P31m$",
    158: r"$P3c1$",
    159: r"$P31c$",
    160: r"$R3m$",
    161: r"$R3c$",
    162: r"$P\bar{3}1m$",
    163: r"$P\bar{3}1c$",
    164: r"$P\bar{3}m1$",
    165: r"$P\bar{3}c1$",
    166: r"$R\bar{3}m$",
    167: r"$R\bar{3}c$",

    # Hexagonal
    168: r"$P6$",
    169: r"$P6_1$",
    170: r"$P6_5$",
    171: r"$P6_2$",
    172: r"$P6_4$",
    173: r"$P6_3$",
    174: r"$P\bar{6}$",
    175: r"$P6/m$",
    176: r"$P6_3/m$",
    177: r"$P622$",
    178: r"$P6_{1}22$",
    179: r"$P6_{5}22$",
    180: r"$P6_{2}22$",
    181: r"$P6_{4}22$",
    182: r"$P6_{3}22$",
    183: r"$P6mm$",
    184: r"$P6cc$",
    185: r"$P6_3cm$",
    186: r"$P6_3mc$",
    187: r"$P\bar{6}m2$",
    188: r"$P\bar{6}c2$",
    189: r"$P\bar{6}2m$",
    190: r"$P\bar{6}2c$",
    191: r"$P6/mmm$",
    192: r"$P6/mcc$",
    193: r"$P6_3/mcm$",
    194: r"$P6_3/mmc$",

    # Cubic
    195: r"$P23$",
    196: r"$F23$",
    197: r"$I23$",
    198: r"$P2_{1}3$",
    199: r"$I2_{1}3$",
    200: r"$Pm\bar{3}$",
    201: r"$Pn\bar{3}$",
    202: r"$Fm\bar{3}$",
    203: r"$Fd\bar{3}$",
    204: r"$Im\bar{3}$",
    205: r"$Pa\bar{3}$",
    206: r"$Ia\bar{3}$",
    207: r"$P432$",
    208: r"$P4_{2}32$",
    209: r"$F432$",
    210: r"$F4_{1}32$",
    211: r"$I432$",
    212: r"$P4_{3}32$",
    213: r"$P4_{1}32$",
    214: r"$I4_{1}32$",
    215: r"$P\bar{4}3m$",
    216: r"$F\bar{4}3m$",
    217: r"$I\bar{4}3m$",
    218: r"$P\bar{4}3n$",
    219: r"$F\bar{4}3c$",
    220: r"$I\bar{4}3d$",
    221: r"$Pm\bar{3}m$",
    222: r"$Pn\bar{3}n$",
    223: r"$Pm\bar{3}n$",
    224: r"$Pn\bar{3}m$",
    225: r"$Fm\bar{3}m$",
    226: r"$Fm\bar{3}c$",
    227: r"$Fd\bar{3}m$",
    228: r"$Fd\bar{3}c$",
    229: r"$Im\bar{3}m$",
    230: r"$Ia\bar{3}d$",
}


def latex_to_omml(latex_expr):
    """
    Convert a LaTeX math expression to OMML using pandoc.
    Strips existing $ delimiters before wrapping in display math.
    """
    # Strip any existing $ or $$ delimiters
    expr = latex_expr.strip()
    expr = re.sub(r'^\$\$|\$\$$', '', expr).strip()
    expr = re.sub(r'^\$|\$$', '', expr).strip()

    # Wrap in display math for pandoc
    md_input = f"$$\n{expr}\n$$"

    result = subprocess.run(
        ["pandoc", "-f", "markdown", "-t", "docx", "-o", "temp_eq.docx"],
        input=md_input.encode("utf-8"),
        capture_output=True
    )

    if result.returncode != 0:
        raise RuntimeError(f"Pandoc error: {result.stderr.decode()}")

    # Open the temp docx and extract the OMML equation element
    temp_doc = Document("temp_eq.docx")
    for para in temp_doc.paragraphs:
        for child in para._element:
            if child.tag.endswith("}oMath") or child.tag.endswith("}oMathPara"):
                return child

    raise ValueError(f"No equation found in pandoc output for: {latex_expr}")


def insert_equation_in_cell(cell, latex_expr):
    """Insert a LaTeX equation into a table cell as native Word math."""
    import copy
    
    omml_element = latex_to_omml(latex_expr)
    para = cell.paragraphs[0]._element
    para.append(copy.deepcopy(omml_element))


# --- Build the Word document with a math table ---
# doc = Document()
# doc.add_heading("Equations Table", level=1)

# # Define table data: (label, latex)
# equations = [
#     (str(k), v) for k, v in space_groups.items()
# ]

# print(equations[:3])

# table = doc.add_table(rows=len(equations) + 1, cols=2)
# table.style = "Table Grid"

# # Header row
# hdr = table.rows[0].cells
# hdr[0].text = "Formula Name"
# hdr[1].text = "Equation"

# # Make header bold
# for cell in hdr:
#     for para in cell.paragraphs:
#         for run in para.runs:
#             run.bold = True

# # Fill rows with equations
# for i, (label, latex) in enumerate(equations, start=1):
#     table.cell(i, 0).text = label
#     try:
#         insert_equation_in_cell(table.cell(i, 1), latex)
#         print(f"✓ Inserted: {label}")
#     except Exception as e:
#         table.cell(i, 1).text = f"[Error: {e}]"
#         print(f"✗ Failed: {label} → {e}")

# # Cleanup temp file
# import os
# if os.path.exists("temp_eq.docx"):
#     os.remove("temp_eq.docx")

# doc.save("math_table.docx")
# print("\n✅ Saved: math_table.docx")